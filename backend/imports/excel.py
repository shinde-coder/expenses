from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from catalogs.lookups import find_category, find_payment_method, resolve_relationship
from families.models import FamilyMember
from imports.models import ImportBatch, ImportRowError
from ledger.models import FinancialTransaction
from planning.models import CategoryBudget, MonthlyBudget


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value


def _json_safe(raw):
    safe = {}
    for key, value in raw.items():
        if value is None:
            safe[key] = ""
        elif isinstance(value, (datetime, date, Decimal)):
            safe[key] = str(value)
        elif isinstance(value, (str, int, float, bool)):
            safe[key] = value
        else:
            safe[key] = str(value)
    return safe


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    raise ValueError("invalid_date")


def _parse_amount(value):
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, TypeError):
        raise ValueError("invalid_amount")
    if amount <= 0:
        raise ValueError("invalid_amount")
    return amount


def _get_or_create_member(family, paid_by):
    name = paid_by.strip()
    relationship = resolve_relationship(name)
    member = FamilyMember.objects.filter(
        family=family, display_name__iexact=name, is_active=True
    ).first()
    if member:
        return member
    if relationship == FamilyMember.Relationship.SELF:
        member = FamilyMember.objects.filter(
            family=family, relationship=FamilyMember.Relationship.SELF, is_active=True
        ).first()
        if member:
            return member
    return FamilyMember.objects.create(
        family=family, display_name=name, relationship=relationship
    )


def import_workbook(path, *, family, user, dry_run=True, filename=None, lenient=False):
    batch = ImportBatch.objects.create(
        family=family,
        created_by=user,
        filename=filename or str(path),
        dry_run=dry_run,
        status=ImportBatch.Status.PENDING,
    )
    workbook = load_workbook(path, data_only=True)
    if "Transactions" not in workbook.sheetnames:
        batch.status = ImportBatch.Status.FAILED
        batch.summary = {"error": "Transactions sheet not found."}
        batch.save()
        return batch

    sheet = workbook["Transactions"]
    success = failed = duplicates = 0
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        raw = {
            "date": _cell(row[0]) if row else "",
            "category": _cell(row[2]) if row and len(row) > 2 else "",
            "subcategory": _cell(row[3]) if row and len(row) > 3 else "",
            "description": _cell(row[4]) if row and len(row) > 4 else "",
            "payment_method": _cell(row[5]) if row and len(row) > 5 else "",
            "paid_by": _cell(row[6]) if row and len(row) > 6 else "",
            "amount": _cell(row[7]) if row and len(row) > 7 else "",
            "notes": _cell(row[8]) if row and len(row) > 8 else "",
        }
        if not any(raw.values()):
            continue
        try:
            occurred_on = _parse_date(raw["date"] or None)
            amount = _parse_amount(raw["amount"])
            category_name = str(raw["category"]).strip()
            if not category_name:
                raise ValueError("invalid_category")
            category = find_category(family, category_name)
            if category is None:
                raise ValueError("invalid_category")
            subcategory = None
            sub_name = str(raw["subcategory"]).strip() if raw["subcategory"] else ""
            notes = str(raw["notes"] or "").strip()
            if sub_name:
                subcategory = category.subcategories.filter(name__iexact=sub_name).first()
                if subcategory is None:
                    if not lenient:
                        raise ValueError("invalid_subcategory")
                    notes = _append_note(
                        notes, f"Imported without subcategory '{sub_name}' (not under {category.name})."
                    )
            method_name = str(raw["payment_method"]).strip()
            method = find_payment_method(family, method_name)
            if method is None:
                raise ValueError("invalid_payment_method")
            paid_by = str(raw["paid_by"]).strip()
            if not paid_by:
                if not lenient:
                    raise ValueError("invalid_member")
                paid_by = "Me"
                notes = _append_note(notes, "Paid By was empty; assigned to Me.")
            description = str(raw["description"] or "").strip()
            duplicate = FinancialTransaction.objects.filter(
                family=family,
                is_deleted=False,
                occurred_on=occurred_on,
                amount=amount,
                category=category,
                description=description,
            ).exists()
            if duplicate:
                duplicates += 1
                ImportRowError.objects.create(
                    batch=batch,
                    row_number=index,
                    code="duplicate",
                    reason="A matching transaction already exists.",
                    raw=_json_safe(raw),
                )
                continue
            if not dry_run:
                member = _get_or_create_member(family, paid_by)
                FinancialTransaction.objects.create(
                    family=family,
                    created_by=user,
                    type=category.type,
                    amount=amount,
                    occurred_on=occurred_on,
                    category=category,
                    subcategory=subcategory,
                    payment_method=method,
                    member=member,
                    description=description,
                    notes=notes,
                )
            success += 1
        except ValueError as exc:
            failed += 1
            code = str(exc)
            reasons = {
                "invalid_date": "Invalid or missing date.",
                "invalid_amount": "Amount must be a number greater than zero.",
                "invalid_category": "Unknown category.",
                "invalid_subcategory": "Subcategory does not belong to the selected category.",
                "invalid_payment_method": "Unknown payment method.",
                "invalid_member": "Paid By is required.",
            }
            ImportRowError.objects.create(
                batch=batch,
                row_number=index,
                code=code,
                reason=reasons.get(code, code),
                raw=_json_safe(raw),
            )

    batch.success_count = success
    batch.failed_count = failed
    batch.duplicate_count = duplicates
    batch.status = ImportBatch.Status.COMPLETED
    batch.summary = {
        "successful_records": success,
        "failed_records": failed,
        "duplicate_records": duplicates,
    }
    if not dry_run:
        import_budgets(path, family=family)
    batch.save()
    return batch


def _append_note(notes, extra):
    return f"{notes} {extra}".strip() if notes else extra


def import_budgets(path, *, family):
    workbook = load_workbook(path, data_only=True)
    if "Budget" not in workbook.sheetnames:
        return
    month = date(2026, 9, 1)
    if "Summary" in workbook.sheetnames:
        selected = workbook["Summary"]["B4"].value
        if isinstance(selected, datetime):
            month = selected.date().replace(day=1)
        elif isinstance(selected, date):
            month = selected.replace(day=1)
    sheet = workbook["Budget"]
    overall = None
    for row in sheet.iter_rows(min_row=5, values_only=True):
        name = _cell(row[0]) if row else ""
        amount = row[1] if row and len(row) > 1 else None
        if not name or amount is None:
            continue
        if str(name).strip().upper() == "TOTAL":
            try:
                overall = Decimal(str(amount))
            except InvalidOperation:
                overall = None
            continue
        category = find_category(family, str(name).strip())
        if category is None:
            continue
        try:
            budget_amount = Decimal(str(amount))
        except InvalidOperation:
            continue
        CategoryBudget.objects.update_or_create(
            family=family,
            year_month=month,
            category=category,
            defaults={"amount": budget_amount},
        )
    MonthlyBudget.objects.update_or_create(
        family=family,
        year_month=month,
        defaults={"overall_amount": overall},
    )


def batch_payload(batch):
    return {
        "id": str(batch.id),
        "filename": batch.filename,
        "dry_run": batch.dry_run,
        "status": batch.status,
        "successful_records": batch.success_count,
        "failed_records": batch.failed_count,
        "duplicate_records": batch.duplicate_count,
        "errors": [
            {
                "row": err.row_number,
                "code": err.code,
                "reason": err.reason,
                "raw": err.raw,
            }
            for err in batch.row_errors.all()
        ],
    }
